# ORANLAB Android - offline Excel analysis prototype
# Works locally; no network requests or API integration.
import os, re, threading
from collections import Counter
from kivy.app import App
from kivy.clock import Clock
from kivy.metrics import dp
from kivy.uix.boxlayout import BoxLayout
from kivy.uix.button import Button
from kivy.uix.label import Label
from kivy.uix.textinput import TextInput
from kivy.uix.spinner import Spinner
from kivy.uix.scrollview import ScrollView
from kivy.uix.gridlayout import GridLayout
from kivy.uix.popup import Popup

try:
    from plyer import filechooser
except Exception:
    filechooser = None

MATCH_TERMS = ['tarih','date','saat','time','lig','league','ülke','ulke','country','sezon','season','ms','iy','ft','ht','ev sahibi','evsahibi','home team','hometeam','home','deplasman','away team','awayteam','away','takım','takim','team','maç','mac','match','fixture','skor','score','sonuç','sonuc','oran','odds','bet','bahis','1x2','h2h','btts','kg var','kg yok','alt/üst','alt ust','over','under','totals','gol','goal','handicap','korner','corner','asian','double chance','çifte şans','cifte sans','draw no bet','bookmaker','bookie','closing','opening']

def is_relevant(name, values):
    n=str(name).strip().lower()
    if any(t in n for t in MATCH_TERMS): return True
    nums=[]
    for x in values[:200]:
        try:
            z=float(str(x).replace(',','.'))
            if 1.01 <= z <= 40: nums.append(z)
        except Exception: pass
    return len(nums)>=10 and sum(abs(x-round(x))>0.005 for x in nums)/len(nums)>=.35

def _iter_source(path):
    """Yield rows from supported files without materializing the full workbook."""
    ext=os.path.splitext(path)[1].lower()
    if ext=='.xlsb':
        from pyxlsb import open_workbook
        with open_workbook(path) as wb:
            names=wb.sheets
            sheet=next((s for s in names if s.strip().lower() in ('arsiv','arşiv','data','matches','sheet1')), names[0])
            with wb.get_sheet(sheet) as sh:
                for i,row in enumerate(sh.rows()):
                    yield [c.v for c in row]
                    if i>=39999: break
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        wb=load_workbook(path, read_only=True, data_only=True)
        try:
            ws=next((s for s in wb.sheetnames if s.strip().lower() in ('arsiv','arşiv','data','matches','sheet1')),wb.sheetnames[0])
            for i,row in enumerate(wb[ws].iter_rows(values_only=True)):
                yield list(row)
                if i>=39999: break
        finally:
            wb.close()
    elif ext=='.csv':
        import csv
        with open(path,encoding='utf-8-sig',newline='') as f:
            for i,row in enumerate(csv.reader(f)):
                yield row
                if i>=39999: break
    else:
        raise ValueError('Desteklenen biçimler: .xlsb, .xlsx ve .csv. .xls dosyasını önce .xlsx olarak kaydedin.')

def read_excel(path):
    """Read only match/odds columns to reduce Android memory use."""
    ext=os.path.splitext(path)[1].lower()
    sheet='CSV' if ext=='.csv' else 'Arsiv / otomatik seçilen sayfa'
    # First pass: find the header and sample enough rows to identify likely match/odds fields.
    sample=[]
    for i,row in enumerate(_iter_source(path)):
        sample.append(row)
        if i>=200: break
    if not sample: raise ValueError('Dosyada veri bulunamadı.')
    header_idx=max(range(min(4,len(sample))), key=lambda i: sum(bool(v is not None and str(v).strip()) for v in sample[i]))
    raw_headers=sample[header_idx]
    headers=[str(x).strip() if x is not None and str(x).strip() else f'Sütun_{j+1}' for j,x in enumerate(raw_headers)]
    # Duplicate Excel headings must be unique internally, but remain readable.
    seen={}; unique=[]
    for h in headers:
        seen[h]=seen.get(h,0)+1
        unique.append(h if seen[h]==1 else f'{h} ({seen[h]})')
    headers=unique
    relevant_idx=[]
    for j,h in enumerate(headers):
        vals=[r[j] for r in sample[header_idx+1:] if j<len(r)]
        if is_relevant(h,vals): relevant_idx.append(j)
    # Always retain likely full-time and half-time score columns for outcome analysis.
    for j,h in enumerate(headers):
        hn=h.strip().upper().replace('_',' ')
        if hn in ('MS','İY','IY','MS SKOR','İY SKOR','FT SCORE','HT SCORE','FULL TIME','HALF TIME') and j not in relevant_idx:
            relevant_idx.append(j)
    if not relevant_idx:
        relevant_idx=list(range(len(headers)))
    relevant_idx.sort()
    selected_headers=[headers[j] for j in relevant_idx]
    # Second pass: retain only selected fields, not all 486 columns per row.
    rows=[]
    for i,row in enumerate(_iter_source(path)):
        if i<=header_idx: continue
        values=[]
        has_value=False
        for j in relevant_idx:
            v=row[j] if j<len(row) else None
            values.append(v)
            if v is not None and str(v).strip(): has_value=True
        if has_value:
            rows.append(dict(zip(selected_headers,values)))
    if not rows: raise ValueError('Başlık bulundu ancak altında maç verisi bulunamadı.')
    return selected_headers, rows, sheet

class OranLabAndroid(App):
    def build(self):
        self.title='ORANLAB Android'
        self.rows=[]; self.headers=[]; self.filter_inputs={}; self.filtered=[]
        root=BoxLayout(orientation='vertical',padding=dp(10),spacing=dp(7))
        root.add_widget(Label(text='[b][color=5eead4]ORAN[/color]LAB[/b]  •  ÇEVRİMDIŞI ANALİZ',markup=True,size_hint_y=None,height=dp(38),font_size='20sp'))
        actions=BoxLayout(size_hint_y=None,height=dp(46),spacing=dp(6))
        actions.add_widget(Button(text='Excel / CSV Yükle',on_release=self.pick_file))
        actions.add_widget(Button(text='Filtreleri Uygula',on_release=self.apply_filters))
        actions.add_widget(Button(text='Temizle',on_release=self.clear_filters))
        root.add_widget(actions)
        self.file_label=Label(text='Henüz dosya yüklenmedi',size_hint_y=None,height=dp(28),halign='left',valign='middle')
        self.file_label.bind(size=lambda *_: setattr(self.file_label,'text_size',self.file_label.size))
        root.add_widget(self.file_label)
        self.mode=Spinner(text='MS 1-X-2',values=('MS 1-X-2','İY 1-X-2','Karşılıklı Gol','2.5 Alt/Üst','En sık MS skoru'),size_hint_y=None,height=dp(42))
        self.mode.bind(text=lambda *_: self.update_stats(self.filtered or self.rows))
        root.add_widget(self.mode)
        self.stats=Label(text='Yüzde analizi için Excel yükle.',size_hint_y=None,height=dp(60),halign='left',valign='middle')
        self.stats.bind(size=lambda *_: setattr(self.stats,'text_size',(self.stats.width, None)))
        root.add_widget(self.stats)
        self.status=Label(text='İnternet bağlantısı kullanılmaz.',size_hint_y=None,height=dp(28),halign='left')
        root.add_widget(self.status)
        self.scroll=ScrollView(do_scroll_x=False)
        self.grid=GridLayout(cols=1,spacing=dp(5),size_hint_y=None,padding=(0,0,0,dp(8)))
        self.grid.bind(minimum_height=self.grid.setter('height'))
        self.scroll.add_widget(self.grid); root.add_widget(self.scroll)
        return root

    def pick_file(self,*_):
        if filechooser is None:
            self.notice('Dosya seçici açılamadı. Android dosya seçici desteği kontrol edilmeli.'); return
        try:
            filechooser.open_file(on_selection=self.selected, filters=['*.xlsb','*.xlsx','*.csv'])
        except Exception as e: self.notice(str(e))

    def selected(self, selection):
        if not selection: return
        path=selection[0]
        self.file_label.text='Dosya okunuyor: '+os.path.basename(path)
        self.status.text='Excel okunuyor; büyük arşivlerde biraz sürebilir.'
        threading.Thread(target=self._load_worker,args=(path,),daemon=True).start()

    def _load_worker(self,path):
        try:
            headers,rows,sheet=read_excel(path)
            relevant=list(headers)
            Clock.schedule_once(lambda dt:self.loaded(path,headers,rows,relevant,sheet))
        except Exception as e:
            Clock.schedule_once(lambda dt:self.notice('Excel yükleme hatası: '+str(e)))

    def loaded(self,path,headers,rows,relevant,sheet):
        self.headers=headers; self.rows=rows; self.filtered=rows
        self.file_label.text=f'{os.path.basename(path)} • {len(rows):,} satır • {len(headers)} sütun'
        self.grid.clear_widgets(); self.filter_inputs={}
        self.grid.cols=1
        for col in relevant:
            line=BoxLayout(size_hint_y=None,height=dp(42),spacing=dp(5))
            line.add_widget(Label(text=col,size_hint_x=.48,halign='left',valign='middle',font_size='12sp'))
            inp=TextInput(hint_text='Filtre: metin, !hariç, >1.8, 1.5..2.0',multiline=False,size_hint_x=.52)
            line.add_widget(inp); self.filter_inputs[col]=inp; self.grid.add_widget(line)
        self.status.text=f'{len(relevant)} maç/oran sütunu tanındı • {sheet} • Çevrimdışı'
        self.update_stats(rows)
        self.notice(f'Yüklendi: {len(rows):,} satır. Otomatik filtre sayısı: {len(relevant)}')

    def apply_filters(self,*_):
        if not self.rows: self.notice('Önce Excel yükle.'); return
        active=[(c,w.text.strip()) for c,w in self.filter_inputs.items() if w.text.strip()]
        out=[]
        for row in self.rows:
            ok=True
            for col,q in active:
                v=row.get(col); s='' if v is None else str(v)
                try:
                    if '..' in q:
                        lo,hi=[float(x.strip().replace(',','.')) for x in q.split('..',1)]; z=float(s.replace(',','.')); ok &= lo<=z<=hi
                    elif q.startswith('>='): ok &= float(s.replace(',','.'))>=float(q[2:].replace(',','.'))
                    elif q.startswith('<='): ok &= float(s.replace(',','.'))<=float(q[2:].replace(',','.'))
                    elif q.startswith('>'): ok &= float(s.replace(',','.'))>float(q[1:].replace(',','.'))
                    elif q.startswith('<'): ok &= float(s.replace(',','.'))<float(q[1:].replace(',','.'))
                    elif q.startswith('!'): ok &= q[1:].lower() not in s.lower()
                    else: ok &= q.lower() in s.lower()
                except Exception:
                    if q.startswith(('>','<')) or '..' in q: ok=False
                    elif q.startswith('!'): ok &= q[1:].lower() not in s.lower()
                    else: ok &= q.lower() in s.lower()
                if not ok: break
            if ok: out.append(row)
        self.filtered=out; self.status.text=f'Filtre sonucu: {len(out):,} / {len(self.rows):,} maç • {len(active)} aktif filtre'
        self.update_stats(out)

    def clear_filters(self,*_):
        for w in self.filter_inputs.values(): w.text=''
        self.filtered=self.rows; self.status.text='Filtreler temizlendi.'; self.update_stats(self.rows)

    def update_stats(self,rows):
        if not rows: self.stats.text='Filtre sonucunda maç bulunamadı.'; return
        mode=self.mode.text
        wanted=('İY','IY','İLK YARI','ILK YARI','İY SKOR') if mode=='İY 1-X-2' else ('MS','MS SKOR','MAÇ SONU','MAC SONU','FT SCORE','FULL TIME')
        if mode=='Karşılıklı Gol' or mode=='2.5 Alt/Üst': wanted=('MS','MS SKOR','MAÇ SONU','MAC SONU','FT SCORE','FULL TIME')
        normalized={str(h).strip().upper().replace('_',' '):h for h in self.headers}
        col=next((normalized[k] for k in wanted if k in normalized),None)
        if not col: self.stats.text=f'{mode}: uygun skor sütunu bulunamadı.'; return
        scores=[]
        for r in rows:
            m=re.match(r'^\s*(\d+)\s*[-:]\s*(\d+)\s*$',str(r.get(col,'')))
            if m: scores.append((int(m.group(1)),int(m.group(2))))
        if not scores: self.stats.text=f'{mode}: okunabilir skor bulunamadı.'; return
        n=len(scores); home=sum(a>b for a,b in scores); draw=sum(a==b for a,b in scores); away=n-home-draw
        if mode=='En sık MS skoru':
            counts=Counter(f'{a}-{b}' for a,b in scores); score,count=counts.most_common(1)[0]; self.stats.text=f'En sık skor: {score} • %{count/n*100:.1f} ({count}/{n})'
        elif mode=='Karşılıklı Gol':
            yes=sum(a>0 and b>0 for a,b in scores); self.stats.text=f'KG Var %{yes/n*100:.1f} ({yes}) • KG Yok %{(n-yes)/n*100:.1f} ({n-yes}) • Örneklem {n}'
        elif mode=='2.5 Alt/Üst':
            over=sum(a+b>=3 for a,b in scores); self.stats.text=f'2.5 Üst %{over/n*100:.1f} ({over}) • 2.5 Alt %{(n-over)/n*100:.1f} ({n-over}) • Örneklem {n}'
        else:
            lab='İY' if mode=='İY 1-X-2' else 'MS'; self.stats.text=f'{lab} 1 %{home/n*100:.1f} ({home}) • X %{draw/n*100:.1f} ({draw}) • 2 %{away/n*100:.1f} ({away}) • Örneklem {n}'

    def notice(self,msg):
        p=Popup(title='ORANLAB',content=Label(text=msg),size_hint=(.9,.4)); p.open()

if __name__=='__main__': OranLabAndroid().run()
